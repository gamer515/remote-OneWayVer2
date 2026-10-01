from pathlib import Path

import pdfplumber
from docx import Document


PDFS = [
    Path(r"C:\Users\유영선\Downloads\lecture1- AIML응용프로젝트2_소개.pdf"),
    Path(r"C:\Users\유영선\Downloads\프로젝트_보고서양식.pdf"),
    Path(r"C:\Users\유영선\Downloads\프로젝트_평가표양식_2026.pdf"),
]
DOCX = Path(r"C:\Users\유영선\Desktop\제안서.docx")


for path in PDFS:
    print(f"\n===== PDF {path.name} =====")
    with pdfplumber.open(path) as pdf:
        print(f"pages={len(pdf.pages)}")
        for i, page in enumerate(pdf.pages, 1):
            text = page.extract_text(x_tolerance=2, y_tolerance=2) or ""
            print(f"\n--- page {i} ---\n{text}")

print(f"\n===== DOCX {DOCX.name} =====")
doc = Document(DOCX)
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)}")
for i, paragraph in enumerate(doc.paragraphs):
    if paragraph.text.strip():
        print(f"P{i} [{paragraph.style.name}] {paragraph.text}")
for ti, table in enumerate(doc.tables):
    print(f"\n--- table {ti} rows={len(table.rows)} cols={len(table.columns)} ---")
    for ri, row in enumerate(table.rows):
        cells = [cell.text.replace("\n", " / ").strip() for cell in row.cells]
        print(f"R{ri}: " + " || ".join(cells))
