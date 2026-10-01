from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


DOCX = Path(r"C:\remote-OneWayVer2\outputs\제안서_유영선_검토본.docx")
FONT = "맑은 고딕"


def set_font(run, size=9.0, bold=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    for key in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        rpr.rFonts.set(qn(key), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold


def remove_numbering(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    numpr = ppr.find(qn("w:numPr"))
    if numpr is not None:
        ppr.remove(numpr)


def make_bullet(paragraph):
    text = paragraph.text.strip()
    remove_numbering(paragraph)
    paragraph.clear()
    run = paragraph.add_run(f"• {text}")
    set_font(run, size=10.0)
    paragraph.paragraph_format.left_indent = Pt(18)
    paragraph.paragraph_format.first_line_indent = Pt(-10)


def make_subheading(paragraph):
    remove_numbering(paragraph)
    text = paragraph.text.strip()
    paragraph.clear()
    run = paragraph.add_run(text)
    set_font(run, size=10.5, bold=True)
    paragraph.paragraph_format.left_indent = Pt(18)
    paragraph.paragraph_format.first_line_indent = Pt(0)


def set_page_numbering(doc):
    for index, section in enumerate(doc.sections):
        sectpr = section._sectPr
        pgn = sectpr.find(qn("w:pgNumType"))
        if index == 0:
            if pgn is None:
                pgn = OxmlElement("w:pgNumType")
                sectpr.append(pgn)
            pgn.set(qn("w:start"), "1")
        elif pgn is not None:
            sectpr.remove(pgn)

        section.footer.is_linked_to_previous = False
        footer = section.footer
        for child in list(footer._element):
            footer._element.remove(child)
        paragraph = footer.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)

        left = paragraph.add_run("- ")
        set_font(left, size=9.0)

        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        field_run = OxmlElement("w:r")
        field_rpr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), FONT)
        fonts.set(qn("w:hAnsi"), FONT)
        fonts.set(qn("w:eastAsia"), FONT)
        field_rpr.append(fonts)
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), "18")
        field_rpr.append(size)
        field_run.append(field_rpr)
        text = OxmlElement("w:t")
        text.text = "1"
        field_run.append(text)
        field.append(field_run)
        paragraph._p.append(field)

        right = paragraph.add_run(" -")
        set_font(right, size=9.0)


doc = Document(DOCX)
p = doc.paragraphs

for index in range(64, 70):
    make_bullet(p[index])
for index in range(91, 97):
    make_bullet(p[index])
for index in (75, 82, 86, 90):
    make_subheading(p[index])

set_page_numbering(doc)
doc.save(DOCX)
print(DOCX)
