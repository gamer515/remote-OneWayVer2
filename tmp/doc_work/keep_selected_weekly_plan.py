from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


DOCX = Path(r"C:\remote-OneWayVer2\outputs\제안서_유영선_검토본.docx")
FONT = "맑은 고딕"


def set_run_font(run, size=8.0, bold=False):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    for key in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        rpr.rFonts.set(qn(key), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_cell_text(cell, text, *, align=WD_ALIGN_PARAGRAPH.LEFT, size=8.0):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    set_run_font(run, size=size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_week_row(row, label, completed=(), planned=()):
    set_cell_text(row.cells[0], label, size=7.5)
    for week in range(1, 16):
        value = "D" if week in completed else "P" if week in planned else ""
        set_cell_text(row.cells[week], value, align=WD_ALIGN_PARAGRAPH.CENTER, size=8.0)


doc = Document(DOCX)
plan = doc.tables[1]

selected_rows = [
    ("유영선 요구사항 및 기존 구조 분석", (1, 2), ()),
    ("유영선 선택형 스토리와 저장 구조 정비", (2, 3, 4), ()),
    ("유영선 LLM 이야기 변형 및 응답 검증", (3, 4), (5, 6, 7)),
    ("유영선 terrain prefab과 JSON 환경 구성", (3, 4), (5, 6, 7, 8)),
    ("유영선 코인 및 튜토리얼 미니게임 구현", (4,), (5, 6, 7, 8)),
    ("유영선 Excel JSON 콘텐츠 관리 절차", (), (5, 6, 7, 8, 9, 10)),
    ("유영선 terrain별 환경 이야기 미니게임", (), (6, 7, 8, 9, 10, 11, 12)),
    ("유영선 한국어 LLM 비교 및 파인튜닝", (), (8, 9, 10, 11, 12, 13)),
    ("유영선 통합 성능 사용성 테스트", (), (10, 11, 12, 13, 14)),
]

for row, (label, completed, planned) in zip(plan.rows[1:10], selected_rows):
    set_week_row(row, label, completed, planned)

for row in plan.rows[10:]:
    for cell in row.cells:
        set_cell_text(cell, "", align=WD_ALIGN_PARAGRAPH.CENTER, size=8.0)

doc.save(DOCX)
print(DOCX)
