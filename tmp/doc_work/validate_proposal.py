from pathlib import Path

from docx import Document


path = Path(r"C:\remote-OneWayVer2\outputs\제안서_유영선_검토본.docx")
doc = Document(path)
all_text = "\n".join(
    [paragraph.text for paragraph in doc.paragraphs]
    + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
)

required = [
    "C211201 유영선",
    "Excel과 JSON",
    "terrain별 환경",
    "코인 베팅 및 인벤토리 구조",
    "한국어 중심 LLM",
    "Codex, CLI, MCP",
    "시험평가",
    "주간별/팀원별 과제 수행 계획",
    "게임 내부에 한국어 중심 LLM을 직접 내장",
    "api 통신을 통한 외부 LLM model 연결, 전투 씬의 전투 패턴 설계 및 구현",
    "상태 기반 전투 제어, 코루틴 기반 순차 실행 및 동적 BattleBox 변형 기술",
    "1. 시나리오 일관성 테스트: 다회차 플레이를 통해 LLM이 생성하는 이야기가 맥락에 맞는지 검토",
    "A. api통신을 통해 외부 LLM에게 기존 이야기를 전달하고, 새로운 이야기를 받아야함",
    "B. 전투 상황 발생 시 실시간 탄막 피하기 미니게임이 정상적으로 구동되어야 함.",
    "6. 데이터 계층: 피해량과 난이도 수치, 전투별 흐름 정보 저장",
    "3. 성능 평가:  api통신을 통한 LLM 생성 시간 측정 및 비동기 처리 안정성 확인.",
]
for phrase in required:
    if phrase not in all_text:
        raise SystemExit(f"필수 문구 누락: {phrase}")

for placeholder in ("녹색/파란색", "작성자제외", "팀원3", "다음과 같은 네 단계"):
    if placeholder in all_text:
        raise SystemExit(f"안내 또는 자리표시자 잔존: {placeholder}")

for unwanted in (
    "외부 LLM 모델 연결 보조",
    "B1, B2, B3의 단계 순서, 입력 방식",
):
    if unwanted in all_text:
        raise SystemExit(f"축약 문구 잔존: {unwanted}")

plan_text = "\n".join(
    cell.text for row in doc.tables[1].rows for cell in row.cells
)
junho_tasks = (
    "전투 씬 전투 패턴 설계 및 구현",
    "LLM 모델 탐색",
    "각 LLM 모델 테스트",
    "전투 씬 에셋 생성",
    "전투 씬 호출 방식 변경",
    "수치에 따른 난이도 조정 구현",
)
yys_tasks = (
    "유영선 요구사항 및 기존 구조 분석",
    "유영선 선택형 스토리와 저장 구조 정비",
    "유영선 LLM 이야기 변형 및 응답 검증",
    "유영선 terrain prefab과 JSON 환경 구성",
    "유영선 코인 및 튜토리얼 미니게임 구현",
    "유영선 Excel JSON 콘텐츠 관리 절차",
    "유영선 terrain별 환경 이야기 미니게임",
    "유영선 한국어 LLM 비교 및 파인튜닝",
    "유영선 통합 성능 사용성 테스트",
)
for selected_task in junho_tasks + yys_tasks:
    if selected_task not in plan_text:
        raise SystemExit(f"일정 누락: {selected_task}")

if max(plan_text.index(task) for task in junho_tasks) > min(plan_text.index(task) for task in yys_tasks):
    raise SystemExit("황준호 일정이 유영선 일정 위에 배치되지 않음")

if len(doc.tables[1].rows) != 16:
    raise SystemExit(f"주간 계획 행 개수 이상: {len(doc.tables[1].rows)}")

if len(doc.tables) != 3:
    raise SystemExit(f"표 개수 이상: {len(doc.tables)}")
if len(doc.sections) != 2:
    raise SystemExit(f"구역 개수 이상: {len(doc.sections)}")

print(f"검증 완료: paragraphs={len(doc.paragraphs)}, tables={len(doc.tables)}, sections={len(doc.sections)}")
