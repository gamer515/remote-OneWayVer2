from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


DOCX = Path(r"C:\remote-OneWayVer2\outputs\제안서_유영선_검토본.docx")
FONT = "맑은 고딕"


def set_run_font(run, size=None, bold=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    for key in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        rpr.rFonts.set(qn(key), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def clear_direct_numbering(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    numpr = ppr.find(qn("w:numPr"))
    if numpr is not None:
        ppr.remove(numpr)


def replace(paragraph, text, *, bold=False, size=None, indent=0, first=0):
    clear_direct_numbering(paragraph)
    paragraph.clear()
    paragraph.paragraph_format.left_indent = Pt(indent)
    paragraph.paragraph_format.first_line_indent = Pt(first)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)


def replace_labeled(paragraph, label, text):
    clear_direct_numbering(paragraph)
    paragraph.clear()
    paragraph.paragraph_format.left_indent = Pt(0)
    paragraph.paragraph_format.first_line_indent = Pt(0)
    label_run = paragraph.add_run(label)
    set_run_font(label_run, bold=True)
    body_run = paragraph.add_run(text)
    set_run_font(body_run)


def insert_before(anchor, text, *, kind="body"):
    style = "Body Text" if kind in ("body", "subheading") else "List Paragraph"
    paragraph = anchor.insert_paragraph_before(style=style)
    if kind == "subheading":
        replace(paragraph, text, bold=True, size=10.5, indent=18)
    elif kind == "bullet":
        replace(paragraph, f"• {text}", size=10.0, indent=28, first=-10)
    else:
        replace(paragraph, text, size=10.0, indent=18)
    return paragraph


def set_cell_text(cell, text, *, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=8.0):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_week_row(row, label, values):
    set_cell_text(row.cells[0], label, size=7.5)
    for week in range(1, 16):
        set_cell_text(
            row.cells[week],
            values.get(week, ""),
            align=WD_ALIGN_PARAGRAPH.CENTER,
            size=8.0,
        )


doc = Document(DOCX)
p = doc.paragraphs

# 요소 기술에서 유영선 담당을 구체적으로 설명하고 황준호 담당은 원본으로 복구한다.
replace_labeled(
    p[40],
    "유영선: ",
    "게임 내부에 한국어 중심 LLM을 직접 내장하고, 플레이 중 축적되는 선택과 사건 기록을 이용해 다음 이야기의 표현을 바꾸는 동적 서사 환경을 설계한다. 네트워크 연결이 없어도 이야기 변경 기능을 사용할 수 있도록 모델 크기, 메모리 사용량, 추론 속도를 비교하고, 게임에 포함할 수 있는 수준으로 경량화한다. 또한 terrain마다 환경 콘셉트, 등장 인물, 사건, 선택지와 미니게임을 연결하여 이야기 변화가 실제 플레이 콘텐츠로 이어지도록 구성한다.",
)
replace(p[41], "황준호:", bold=True)
replace(p[42], "Api 기반 LLM Model")
replace(
    p[43],
    "기존에 사용하던 gemini2.5flash를 비롯하여 Llama 3, Gemma 2, Mixtral, Qwen 2.5, Phi-4등 무료로 지원하는 다양한 모델을 사용하여 1학기에 있던 적은 토큰 문제를 보완",
)

replace_labeled(
    p[47],
    "유영선: ",
    "현재 구현된 선택형 스토리 UI, Story JSON 구조, 플레이 기록과 세이브 데이터, 코인 베팅 및 인벤토리 구조를 계속 발전시키고 있다. 기존 기능을 단순히 재사용하는 데 그치지 않고 terrain별 환경 세부 구성, Encounter별 이야기와 선택 조건, 해금 요소, 미니게임 진입과 결과 반영을 추가하여 콘텐츠의 밀도와 완성도를 높이고 있다.",
)
replace(
    p[48],
    "황준호: Undertale (전투 시스템): 특정 이벤트 발생 시 전환되는 탄막 회피 미니게임 인터페이스 및 조작감 벤치마킹.",
)

replace_labeled(
    p[51],
    "유영선: ",
    "외부 서버에만 의존하지 않고 한국어 이야기에 익숙한 LLM을 게임 내부에 내장한다. 플레이어의 선택 기록, 현재 terrain, 등장 인물의 관계와 해금 상태를 입력으로 구성하고, 모델이 기존 이야기의 핵심 사건을 유지하면서 문장과 세부 전개를 변경하도록 학습시킨다. 생성 결과는 지정된 JSON 구조와 허용된 이야기 범위를 검사한 뒤 적용하여, 반복 플레이의 변화와 게임 진행의 안정성을 함께 확보한다.",
)
replace(
    p[52],
    "황준호: 기존 Undertale 게임에서 새로운 패턴을 추가함, 날라오는 물체 자르기나 패링 등 새로운 조작에 의한 패턴들을 추가",
)

replace_labeled(
    p[54],
    "유영선: ",
    "terrain별 지형과 분위기는 prefab으로 구성하고, NPC, 상호작용 위치, 이야기 연결, 선택지, 해금 조건과 미니게임 설정은 JSON으로 분리한다. Excel에서는 전체 terrain과 사건, 등장 인물, 아이템, 미니게임의 관계를 한 번에 검토하고 수정하며, 검증된 데이터를 JSON으로 변환해 Unity가 읽도록 설계한다. 내장형 LLM은 원문 이야기와 플레이 기록을 입력받아 다음 회차에 사용할 이야기 변경안을 생성하고, 형식과 참조가 정상인 결과만 저장한다. 이를 통해 환경 제작, 데이터 관리, 이야기 변화와 미니게임 추가를 하나의 반복 가능한 콘텐츠 제작 절차로 묶는다.",
)
replace(
    p[55],
    "황준호: 상태 기반 전투 제어, 코루틴 기반 순차 실행 및 동적 BattleBox 변형 기술을 활용하여 탄막 회피 전투를 구현하였다. 또한 B1, B2, B3가 각각 독립적인 진행 순서와 조작 방식을 가질 수 있도록 전투 흐름, 실행 단계 및 개별 탄막 패턴을 분리한 모듈형 구조를 적용한다. 이를 통해 마우스 기반 전진, 날아오는 물체 베기 및 화면 파괴 연출과 같은 새로운 전투 방식을 추가하고, 향후 패턴 확장 시 기존 시스템의 변경 범위를 최소화한다.",
)

# 요구사항: 유영선 담당을 상세히 풀고 황준호 담당 블록은 원문 그대로 유지한다.
replace(p[61], "유영선:", bold=True)
requirements_anchor = p[62]
insert_before(requirements_anchor, "기능적 요구사항", kind="subheading")
insert_before(requirements_anchor, "한국어 중심 LLM을 게임 실행 환경에 내장하여 외부 API 연결 여부와 관계없이 이야기 변경 기능을 실행할 수 있어야 한다.", kind="bullet")
insert_before(requirements_anchor, "플레이어의 선택 기록, 방문한 사건, 획득 아이템, 인물 관계, 해금 상태와 현재 terrain 정보를 모델 입력으로 구성해야 한다.", kind="bullet")
insert_before(requirements_anchor, "각 terrain은 고유한 환경 prefab을 사용하고, NPC와 상호작용 오브젝트는 JSON의 식별자와 위치 정보에 따라 배치되어야 한다.", kind="bullet")
insert_before(requirements_anchor, "Excel에서 terrain, Story, Interaction, 아이템과 미니게임 데이터를 수정한 뒤 참조 오류를 검사하여 JSON 실행 데이터로 반영할 수 있어야 한다.", kind="bullet")
insert_before(requirements_anchor, "이야기 사건에서 미니게임으로 진입하고, 성공 또는 실패 결과가 보상, 해금, 다음 선택지와 후속 이야기에 반영되어야 한다.", kind="bullet")
insert_before(requirements_anchor, "비기능적 요구사항", kind="subheading")
insert_before(requirements_anchor, "내장 모델은 목표 하드웨어의 메모리 범위 안에서 실행되어야 하며, 이야기 진행을 방해하지 않는 응답 시간을 확보해야 한다.", kind="bullet")
insert_before(requirements_anchor, "생성 결과는 인물, 장소, 이전 선택과 모순되지 않아야 하고, 요청한 JSON 필드 외의 데이터와 식별자를 변경해서는 안 된다.", kind="bullet")
insert_before(requirements_anchor, "새로운 terrain과 미니게임은 기존 코드를 대규모로 수정하지 않고 데이터와 모듈 추가 방식으로 연결할 수 있어야 한다.", kind="bullet")

replace(p[62], "황준호:", bold=True)
replace(p[63], "기능적 요구사항:", bold=True)
replace(p[64], "api통신을 통해 외부 LLM에게 기존 이야기를 전달하고, 새로운 이야기를 받아야함")
replace(p[65], "전투 상황 발생 시 실시간 탄막 피하기 미니게임이 정상적으로 구동되어야 함.")
replace(p[66], "비기능적 요구사항:", bold=True)
replace(p[67], "반응성: 대기 시간 내에 LLM API 통신을 완료해야함")
replace(p[68], "일관성:  LLM이 기존 단어를 적절한 단어로 대체 해야됨")
replace(p[69], "전투가 속도감 있게 진행되야함")

# 시스템 설계: 유영선 설계를 황준호 블록과 같은 수준으로 확장하고 황준호 문구는 원본 복구.
replace(p[73], "유영선:", bold=True)
design_anchor = p[74]
insert_before(design_anchor, "전체 시스템 계층", kind="subheading")
insert_before(design_anchor, "콘텐츠 작성 계층: Excel에서 terrain, 등장 인물, Story, Interaction, 아이템, 해금 조건과 미니게임 목록 및 연결 관계를 관리한다.", kind="bullet")
insert_before(design_anchor, "데이터 변환 및 검증 계층: Excel 항목을 JSON 구조에 대응시키고, placeId, prefabId, 이야기 경로, 해금 키의 누락과 중복을 검사한다.", kind="bullet")
insert_before(design_anchor, "환경 구성 계층: terrain prefab에 고정 지형과 장식을 저장하고, JSON에 따라 NPC, 상호작용 오브젝트와 사건 위치를 배치한다.", kind="bullet")
insert_before(design_anchor, "플레이 진행 계층: DecisionManager가 이야기 카드, 선택, 이동, Encounter, 미니게임 진입과 결과 반영 순서를 제어한다.", kind="bullet")
insert_before(design_anchor, "내장형 LLM 계층: 플레이 기록과 현재 게임 상태를 모델 입력으로 만들고, 로컬 추론 결과의 한국어 문맥과 JSON 형식을 검증한다.", kind="bullet")
insert_before(design_anchor, "저장 계층: 플레이 위치, 선택 이력, 인벤토리, 코인, 관계, 해금 상태, 모델 입력과 생성 이야기를 회차별로 저장한다.", kind="bullet")
insert_before(design_anchor, "주요 자료구조", kind="subheading")
insert_before(design_anchor, "Terrain 데이터: terrain 식별자, prefab, 길이, 배치 지점과 연결된 Encounter 목록을 저장한다.", kind="bullet")
insert_before(design_anchor, "Encounter 데이터: placeId, 등장 인물, Story와 Interaction 경로, 요구 해금, 부여 해금과 실행할 미니게임을 저장한다.", kind="bullet")
insert_before(design_anchor, "내장 LLM 입력 데이터: 원문 이야기, 플레이한 카드, 선택 결과, 인물 관계, 현재 terrain, 아이템과 해금 상태를 저장한다.", kind="bullet")
insert_before(design_anchor, "내장 LLM 출력 데이터: 변경 대상 카드, 카드 순번, 수정 문장, 모델 버전과 검증 결과를 저장한다.", kind="bullet")
insert_before(design_anchor, "미니게임 데이터: 실행 조건, 입력 방식, 난이도, 성공 및 실패 결과, 보상과 후속 이야기 연결을 저장한다.", kind="bullet")
insert_before(design_anchor, "주요 알고리즘", kind="subheading")
insert_before(design_anchor, "콘텐츠 변환 알고리즘: Excel 행을 JSON 객체로 변환하고 식별자와 경로의 참조 무결성을 검사한다.", kind="bullet")
insert_before(design_anchor, "terrain 구성 알고리즘: 현재 장과 에피소드에 맞는 prefab을 불러온 뒤 JSON 배치 정보로 NPC와 사건을 등록한다.", kind="bullet")
insert_before(design_anchor, "이야기 입력 구성 알고리즘: 저장된 플레이 기록에서 다음 이야기 변경에 필요한 선택과 사건만 추출하여 내장 LLM 입력으로 묶는다.", kind="bullet")
insert_before(design_anchor, "로컬 이야기 변형 알고리즘: 파인튜닝한 한국어 LLM으로 허용된 카드의 문장을 변경하고, 카드 수, 식별자, 순번, JSON 형식과 금지된 변경 여부를 검사한다.", kind="bullet")
insert_before(design_anchor, "미니게임 연결 알고리즘: Encounter의 행동과 조건에 따라 미니게임을 선택하고 결과를 보상, 해금, 관계 수치와 다음 Story에 반영한다.", kind="bullet")
insert_before(design_anchor, "실패 복구 알고리즘: 모델 실행 실패, 데이터 누락 또는 참조 오류가 발생하면 로그를 남기고 검증된 기본 이야기와 설정으로 계속 진행한다.", kind="bullet")

replace(p[74], "황준호:", bold=True)
replace(p[75], "전체 시스템 계층:", bold=True)
replace(p[76], "표현 계층: 사용자에게 보이는 화면과 피드백을 담당")
replace(p[77], "입력 계층: 현재 전투 단계에 맞는 사용자 입력을 수집")
replace(p[78], "전투 제어 계층: BattleScene 전체의 실행 순서와 생명주기를 관리")
replace(p[79], "패턴 계층: 실제 탄막과 장애물의 행동을 담당")
replace(p[80], "물리 및 판정 계층: BattleBox 경계 제한, 탄막 충돌, 피해 및 무적 시간, 종료 지점 도달 판정 담당")
replace(p[81], "데이터 계층: 피해량과 난이도 수치, 전투별 흐름 정보 저장")
replace(p[82], "씬 컨트롤:", bold=True)
replace(p[83], "BattleScene은 하나의 공통 씬을 사용하되, 진입 시 현재 전투 식별값을 확인하여 B1, B2, B3 중 하나를 실행")
replace(p[84], "대사, 회피, 게이지 및 특수 연출의 실제 순서는 각 BattleFlow가 결정")
replace(p[85], "각 단계는 고정된 대기 시간만으로 연결하지 않고 대사 종료, 애니메이션 종료, 패턴 완료 및 목표 지점 도달 등의 이벤트를 기준으로 다음 단계로 전환")
replace(p[86], "주요 자료구조(데이터:", bold=True)
replace(p[87], "전투 식별 데이터(BattleType): 현재 실행할 전투를 구분")
replace(p[88], "전투 흐름 데이터(BattleFlowDefinition): 실행할 단계 목록, 시작 설정, 성공 조건, 실패 조건 저장")
replace(p[89], "패턴 설정 데이터(BattlePatternData): 패턴 이름, 지속 시간, 탄환 종류, 생성 간격, 탄환 속도, 피해량, BattleBox 크기와 위치, 플레이어 이동 모드, 종료 후 대기 시간")
replace(p[90], "주요 알고리즘:", bold=True)
replace(p[91], "전투 흐름 선택 알고리즘: BattleScene 진입 시 현재 전투 식별값을 확인 및 전투 완료시 씬 전환")
replace(p[92], "탄막 패턴 선택 알고리즘:")
replace(p[93], "순차 실행: 등록된 순서대로 패턴을 실행한다.")
replace(p[94], "무작위 실행: 패턴 목록에서 중복 없이 무작위로 선택한다.")
replace(p[95], "가중치 실행: 난이도와 등장 확률에 따라 패턴을 선택한다.")
replace(p[96], "조건부 실행: 체력, 경과 시간 및 이전 패턴 결과에 따라 다음 패턴을 결정한다.")

# 구현 환경에서 담당 방식을 명확히 분리한다.
replace(
    p[104],
    "AI와 자동화 도구: 유영선은 한국어 중심 LLM의 파인튜닝, 경량화 및 게임 내장형 추론 환경을 구성하고, 황준호는 외부 LLM API 통신을 이용한다. 개발 과정에서는 Codex, CLI, MCP, 플러그인과 skill을 활용한다.",
)

# 시험평가: 유영선 시험 계획을 보강하고 황준호 블록은 원문 복구.
replace(p[110], "유영선:", bold=True)
test_anchor = p[111]
insert_before(test_anchor, "내장 모델 시험: 인터넷 연결 없이 모델이 정상적으로 로드되고 이야기 변경 결과를 반환하는지 확인하며, 모델 크기별 메모리 사용량, 초기 로딩 시간과 평균 및 최대 추론 시간을 측정한다.", kind="bullet")
insert_before(test_anchor, "이야기 품질 시험: 같은 플레이 기록을 모델별로 입력하여 한국어 자연스러움, 인물과 장소의 일관성, 핵심 사건 유지 여부, JSON 형식 준수율과 요청 외 변경 비율을 비교한다.", kind="bullet")
insert_before(test_anchor, "데이터 시험: Excel과 JSON의 항목 대응을 확인하고 prefabId, placeId, Story와 Interaction 경로, 해금 키의 누락, 중복과 잘못된 참조를 검사한다.", kind="bullet")
insert_before(test_anchor, "환경 및 저장 시험: terrain별 prefab과 JSON 배치 결과를 확인하고 재실행 후 위치, 선택 기록, 인벤토리, 코인, 관계, 해금과 생성 이야기가 동일하게 복원되는지 검증한다.", kind="bullet")
insert_before(test_anchor, "미니게임 시험: 이야기에서 각 미니게임으로 정상 진입하는지 확인하고 성공, 실패, 취소 결과가 보상, 해금과 다음 이야기에 올바르게 반영되는지 테스트한다.", kind="bullet")

replace(p[111], "황준호:", bold=True)
replace(p[112], "시나리오 일관성 테스트: 다회차 플레이를 통해 LLM이 생성하는 이야기가 맥락에 맞는지 검토")
replace(p[113], "사용성 테스트: 실제 사용자 대상으로 탄막 전투의 속도감, 난이도, 다양성 검토")
replace(p[114], "성능 평가:  api통신을 통한 LLM 생성 시간 측정 및 비동기 처리 안정성 확인.")

replace_labeled(
    p[118],
    "유영선: ",
    "한국어 중심 LLM을 게임에 내장하기 위한 모델 선정, 이야기 변경용 데이터 구성, 파인튜닝, 경량화, 로컬 추론과 결과 검증을 담당한다. terrain prefab과 Excel 및 JSON 기반 콘텐츠 관리 구조를 설계하고, terrain별 환경 세부 구성, Encounter, 등장 인물, 선택지, 해금 요소, 이야기와 연결되는 미니게임을 추가한다. 선택과 탐색 흐름, 코인 및 인벤토리, 저장과 회차별 이야기 변경 기능을 계속 발전시키고, Codex, CLI, MCP, 플러그인과 skill을 이용해 개발 및 검증 환경을 정리한다.",
)
replace(p[119], "황준호: api 통신을 통한 외부 LLM model 연결, 전투 씬의 전투 패턴 설계 및 구현")

# 주간 계획: 유영선 계획 6개를 추가하고 황준호 계획 6개는 원본 값으로 복구한다.
plan = doc.tables[1]
yys_rows = [
    ("유영선 내장 LLM 후보 조사 및 실행 환경 설계", {1: "D", 2: "D", 3: "D", 4: "P", 5: "P"}),
    ("유영선 한국어 이야기 데이터 구성 및 파인튜닝", {4: "P", 5: "P", 6: "P", 7: "P", 8: "P", 9: "P"}),
    ("유영선 Excel JSON 콘텐츠 관리 및 검증", {4: "P", 5: "P", 6: "P", 7: "P", 8: "P", 9: "P", 10: "P"}),
    ("유영선 terrain별 환경 이야기 구성", {5: "P", 6: "P", 7: "P", 8: "P", 9: "P", 10: "P", 11: "P", 12: "P"}),
    ("유영선 미니게임 설계 구현 및 이야기 연결", {6: "P", 7: "P", 8: "P", 9: "P", 10: "P", 11: "P", 12: "P", 13: "P"}),
    ("유영선 내장 모델 통합 시험 및 최종 개선", {9: "P", 10: "P", 11: "P", 12: "P", 13: "P", 14: "P", 15: "P"}),
]
junho_rows = [
    ("전투 씬 전투 패턴 설계 및 구현", {1: "D", 2: "D", 3: "D", 4: "D", 5: "P", 6: "P", 7: "P", 8: "P", 9: "P", 10: "P", 11: "P", 12: "P", 13: "P"}),
    ("LLM 모델 탐색", {4: "D", 5: "P", 6: "P"}),
    ("각 LLM 모델 테스트", {6: "P", 7: "P", 8: "P", 9: "P", 10: "P"}),
    ("전투 씬 에셋 생성", {1: "D", 6: "p", 11: "P", 15: "P"}),
    ("전투 씬 호출 방식 변경", {3: "D"}),
    ("수치에 따른 난이도 조정 구현", {12: "P", 13: "P", 14: "P"}),
]
for row, (label, values) in zip(plan.rows[1:], yys_rows + junho_rows):
    set_week_row(row, label, values)

doc.save(DOCX)
print(DOCX)
