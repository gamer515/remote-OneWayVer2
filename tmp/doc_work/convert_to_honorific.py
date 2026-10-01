from pathlib import Path
import re
from docx import Document


SOURCE = Path(r"C:\Users\유영선\Desktop\제안서_유영선_검토본.docx")
OUTPUT = Path(r"C:\remote-OneWayVer2\outputs\제안서_유영선_존댓말본.docx")


def replace_suffix(paragraph, old, new):
    if not paragraph.text.endswith(old):
        raise RuntimeError(f"Suffix {old!r} not found in {paragraph.text!r}")
    remaining = len(old)
    for run in reversed(paragraph.runs):
        if remaining == 0:
            break
        take = min(remaining, len(run.text))
        if take:
            run.text = run.text[:-take]
            remaining -= take
    if remaining:
        raise RuntimeError(f"Could not remove suffix {old!r}")
    target = paragraph.runs[-1] if paragraph.runs else paragraph.add_run()
    target.text += new


doc = Document(SOURCE)

# 문장 내부의 서술형 종결을 존댓말로 변환합니다.
run_replacements = [
    ("높이고 있다.", "높이고 있습니다."),
    ("해야 한다.", "해야 합니다."),
    ("안 된다.", "안 됩니다."),
    ("하였다.", "하였습니다."),
    ("한다.", "합니다."),
    ("된다.", "됩니다."),
    ("있다.", "있습니다."),
    ("없다.", "없습니다."),
]
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        for old, new in run_replacements:
            run.text = run.text.replace(old, new)

# 문장형이 아닌 기존 항목은 의미를 유지하면서 정중한 완결문으로 정리합니다.
suffix_rules = {
    "기존에 사용하던 gemini2.5flash를 비롯하여 Llama 3, Gemma 2, Mixtral, Qwen 2.5, Phi-4등 무료로 지원하는 다양한 모델을 사용하여 1학기에 있던 적은 토큰 문제를 보완": ("보완", "보완합니다."),
    "역할: 게임 실행 환경, 씬 전환, 상태 관리, UI, 미니게임 및 전투 로직 구현": ("로직 구현", "로직을 구현합니다."),
    "적용: 선택 중심 인터페이스, 3D terrain 탐색, 코인 기반 상호작용, 탄막 회피 전투를 하나의 플레이 흐름으로 통합": ("통합", "통합합니다."),
    "구조: 기능별 Controller와 데이터 클래스를 분리하여 기능 추가 시 기존 코드의 변경 범위를 줄임": ("줄임", "줄입니다."),
    "역할: 팀 단위 형상 관리, 기능별 작업 분리, 변경 이력과 통합 상태 관리": ("관리", "관리합니다."),
    "적용: 개인 기능을 브랜치와 커밋 단위로 관리하고, 씬 및 prefab 충돌을 확인한 뒤 통합": ("통합", "통합합니다."),
    "역할: 코드베이스 분석, Unity Editor 자동화, 반복 테스트, 에셋 및 문서 제작 보조": ("보조", "보조합니다."),
    "적용: 작업별 지침과 도구를 분리하여 개발 환경을 표준화하고, 생성 결과는 실행 및 렌더링 결과로 검증": ("검증", "검증합니다."),
    "역할: 2D 및 3D 게임 에셋 제작과 편집": ("에셋 제작과 편집", "에셋을 제작하고 편집합니다."),
    "적용: terrain별 시각적 특징, 캐릭터, 상호작용 오브젝트, UI 자산 구성": ("구성", "구성합니다."),
    "황준호: 기존 Undertale 게임에서 새로운 패턴을 추가함, 날라오는 물체 자르기나 패링 등 새로운 조작에 의한 패턴들을 추가": ("추가", "추가합니다."),
    "황준호: Undertale (전투 시스템): 특정 이벤트 발생 시 전환되는 탄막 회피 미니게임 인터페이스 및 조작감 벤치마킹.": ("벤치마킹.", "벤치마킹합니다."),
    "A. api통신을 통해 외부 LLM에게 기존 이야기를 전달하고, 새로운 이야기를 받아야함": ("받아야함", "받아야 합니다."),
    "B. 전투 상황 발생 시 실시간 탄막 피하기 미니게임이 정상적으로 구동되어야 함.": ("되어야 함.", "되어야 합니다."),
    "A. 반응성: 대기 시간 내에 LLM API 통신을 완료해야함": ("완료해야함", "완료해야 합니다."),
    "B. 일관성:  LLM이 기존 단어를 적절한 단어로 대체 해야됨": ("대체 해야됨", "대체해야 합니다."),
    "C. 전투가 속도감 있게 진행되야함": ("진행되야함", "진행되어야 합니다."),
    "1. 표현 계층: 사용자에게 보이는 화면과 피드백을 담당": ("담당", "담당합니다."),
    "2. 입력 계층: 현재 전투 단계에 맞는 사용자 입력을 수집": ("수집", "수집합니다."),
    "3. 전투 제어 계층: BattleScene 전체의 실행 순서와 생명주기를 관리": ("관리", "관리합니다."),
    "4. 패턴 계층: 실제 탄막과 장애물의 행동을 담당": ("담당", "담당합니다."),
    "5. 물리 및 판정 계층: BattleBox 경계 제한, 탄막 충돌, 피해 및 무적 시간, 종료 지점 도달 \t\t판정 담당": ("담당", "담당합니다."),
    "6. 데이터 계층: 피해량과 난이도 수치, 전투별 흐름 정보 저장": ("저장", "저장합니다."),
    "1. BattleScene은 하나의 공통 씬을 사용하되, 진입 시 현재 전투 식별값을 확인하여 B1, B2, B3 \t\t중 하나를 실행": ("실행", "실행합니다."),
    "2. 대사, 회피, 게이지 및 특수 연출의 실제 순서는 각 BattleFlow가 결정": ("결정", "결정합니다."),
    "3. 각 단계는 고정된 대기 시간만으로 연결하지 않고 대사 종료, 애니메이션 종료, 패턴 완료 및 \t\t목표 지점 도달 등의 이벤트를 기준으로 다음 단계로 전환": ("전환", "전환합니다."),
    "1. 전투 식별 데이터(BattleType): 현재 실행할 전투를 구분": ("구분", "구분합니다."),
    "2. 전투 흐름 데이터(BattleFlowDefinition): 실행할 단계 목록, 시작 설정, 성공 조건, 실패 \t\t조건 저장": ("저장", "저장합니다."),
    "3. 패턴 설정 데이터(BattlePatternData): 패턴 이름, 지속 시간, 탄환 종류, 생성 간격, 탄환 \t\t속도, 피해량, BattleBox 크기와 위치, 플레이어 이동 모드, 종료 후 대기 시간": ("종료 후 대기 시간", "종료 후 대기 시간을 저장합니다."),
    "1. 전투 흐름 선택 알고리즘: BattleScene 진입 시 현재 전투 식별값을 확인 및 전투 완료시 씬 \t\t전환": ("씬 \t\t전환", "씬을 전환합니다."),
    "1. 시나리오 일관성 테스트: 다회차 플레이를 통해 LLM이 생성하는 이야기가 맥락에 맞는지 검토": ("검토", "검토합니다."),
    "2. 사용성 테스트: 실제 사용자 대상으로 탄막 전투의 속도감, 난이도, 다양성 검토": ("검토", "검토합니다."),
    "3. 성능 평가:  api통신을 통한 LLM 생성 시간 측정 및 비동기 처리 안정성 확인.": ("확인.", "확인합니다."),
    "황준호: api 통신을 통한 외부 LLM model 연결, 전투 씬의 전투 패턴 설계 및 구현": ("구현", "구현을 담당합니다."),
}

def normalize(text):
    return re.sub(r"\s+", " ", text).strip()


paragraph_by_text = {normalize(p.text): p for p in doc.paragraphs}
missing = []
for full_text, (old, new) in suffix_rules.items():
    paragraph = paragraph_by_text.get(normalize(full_text))
    if paragraph is None:
        missing.append(full_text)
        continue
    replace_suffix(paragraph, old, new)

if missing:
    raise RuntimeError("Missing expected paragraphs:\n" + "\n".join(missing))

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
