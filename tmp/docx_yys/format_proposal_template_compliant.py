from __future__ import annotations

import re
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


SOURCE = Path(r"C:\Users\유영선\Desktop\C211201_유영선.docx")
OUTPUT = Path(r"C:\remote-OneWayVer2\outputs\C211201_유영선_보고서양식준수본.docx")
BODY_FONT = "바탕체"
BODY_SIZE = 10.0
TITLE_SIZE = 12.0


MAJOR_HEADINGS = {
    "과제 개요",
    "필요성 및 동기",
    "요소 기술",
    "상세 내용",
    "주간별/팀원별 과제 수행 계획",
}

SECTION_HEADINGS = {
    "AI 협업 개발 역량과 반복 가능한 제작 과정",
    "기존 게임과 서비스 분석",
    "반복 개발과 통합 검증",
    "공통",
    "기존 기술과 구현 자산의 재사용",
    "차별화기술",
    "중점기술",
    "요구 사항",
    "시스템 설계",
    "구현",
    "시험평가",
    "팀원별 역할",
    "전체 시스템 계층",
    "주요 자료구조",
    "주요 알고리즘",
    "전체 시스템 계층:",
    "씬 컨트롤:",
    "주요 자료구조(데이터:",
    "주요 알고리즘:",
}

SUBHEADINGS = {
    "기능적 요구사항",
    "기능적 요구사항:",
    "비기능적 요구사항",
    "비기능적 요구사항:",
}


def set_run_font(run, size: float = BODY_SIZE, bold: bool | None = None) -> None:
    run.font.name = BODY_FONT
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), BODY_FONT)


def set_common_paragraph_format(paragraph) -> None:
    fmt = paragraph.paragraph_format
    fmt.left_indent = Cm(0)
    fmt.right_indent = Cm(0)
    fmt.first_line_indent = Cm(0)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.3


def format_paragraph(paragraph) -> None:
    text = paragraph.text.strip()
    fmt = paragraph.paragraph_format
    set_common_paragraph_format(paragraph)

    if not text:
        fmt.line_spacing = 1.0
        for run in paragraph.runs:
            set_run_font(run)
        return

    if text in MAJOR_HEADINGS:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.space_before = Pt(9)
        fmt.space_after = Pt(3)
        fmt.line_spacing = 1.3
        for run in paragraph.runs:
            set_run_font(run)
        return

    if text in SECTION_HEADINGS:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.left_indent = Cm(0.4)
        fmt.space_before = Pt(5)
        fmt.space_after = Pt(2)
        fmt.line_spacing = 1.3
        for run in paragraph.runs:
            set_run_font(run)
        return

    if text in SUBHEADINGS:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.left_indent = Cm(0.8)
        fmt.space_before = Pt(3)
        fmt.space_after = Pt(1)
        fmt.line_spacing = 1.3
        for run in paragraph.runs:
            set_run_font(run)
        return

    if text in {"유영선:", "황준호:"}:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.left_indent = Cm(0.4)
        fmt.space_before = Pt(3)
        fmt.space_after = Pt(1)
        fmt.line_spacing = 1.3
        for run in paragraph.runs:
            set_run_font(run)
        return

    if text.startswith("•"):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.left_indent = Cm(1.0)
        fmt.first_line_indent = Cm(-0.45)
        fmt.space_after = Pt(1)
        fmt.line_spacing = 1.3
        for run in paragraph.runs:
            set_run_font(run)
        return

    if re.match(r"^\d+\.\s*", text):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.left_indent = Cm(1.0)
        fmt.first_line_indent = Cm(-0.45)
        fmt.space_after = Pt(1)
        fmt.line_spacing = 1.3
        for run in paragraph.runs:
            set_run_font(run)
        return

    label_prefixes = (
        "역할:", "적용:", "구조:", "운영체제:", "개발 언어:",
        "게임 엔진과 IDE:", "데이터 도구:", "AI와 자동화 도구:",
        "그래픽 도구:", "협업 도구:",
    )
    if text.startswith(label_prefixes):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.left_indent = Cm(0.8)
        fmt.space_after = Pt(1)
        for run in paragraph.runs:
            set_run_font(run)
        return

    if text.startswith(("유영선:", "황준호:")):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        fmt.left_indent = Cm(0.8)
        fmt.space_before = Pt(1)
        fmt.space_after = Pt(2)
        for run in paragraph.runs:
            set_run_font(run)
        return

    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if len(text) >= 35 else WD_ALIGN_PARAGRAPH.LEFT
    fmt.left_indent = Cm(0.4)
    fmt.space_after = Pt(2)
    for run in paragraph.runs:
        set_run_font(run)


def format_table(table, index: int) -> None:
    for row in table.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                fmt = paragraph.paragraph_format
                fmt.left_indent = Cm(0)
                fmt.right_indent = Cm(0)
                fmt.first_line_indent = Cm(0)
                fmt.space_before = Pt(0)
                fmt.space_after = Pt(0)
                fmt.line_spacing = 1.0
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT
                    if index in (1, 2) and c_idx == 0 and r_idx > 0
                    else WD_ALIGN_PARAGRAPH.CENTER
                )

                size = TITLE_SIZE if index == 0 and r_idx == 0 else BODY_SIZE
                for run in paragraph.runs:
                    set_run_font(run, size)


def normalized_text(doc) -> str:
    items: list[str] = []
    items.extend(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                items.extend(p.text for p in cell.paragraphs)
    for section in doc.sections:
        items.extend(p.text for p in section.header.paragraphs)
        items.extend(p.text for p in section.footer.paragraphs)
    return re.sub(r"\s+", " ", "\n".join(items)).strip()


def page_structure(path: Path):
    doc = Document(path)
    sections = [
        (
            str(s.start_type),
            int(s.page_width), int(s.page_height),
            int(s.top_margin), int(s.bottom_margin),
            int(s.left_margin), int(s.right_margin),
            str(s.orientation),
        )
        for s in doc.sections
    ]
    with ZipFile(path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    return {
        "sections": sections,
        "page_breaks": sum(1 for e in root.iter(w + "br") if e.get(w + "type") == "page"),
        "rendered_page_breaks": sum(1 for _ in root.iter(w + "lastRenderedPageBreak")),
        "section_properties": sum(1 for _ in root.iter(w + "sectPr")),
        "section_types": [e.get(w + "val") for e in root.iter(w + "type")],
    }


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(SOURCE)
    before_text = normalized_text(doc)
    before_pages = page_structure(SOURCE)

    for style_name in ("Normal", "Body Text", "List Paragraph"):
        style = doc.styles[style_name]
        style.font.name = BODY_FONT
        style.font.size = Pt(BODY_SIZE)
        rfonts = style._element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            rfonts.set(qn(f"w:{attr}"), BODY_FONT)

    for paragraph in doc.paragraphs:
        format_paragraph(paragraph)

    for index, table in enumerate(doc.tables):
        format_table(table, index)

    doc.save(OUTPUT)
    check = Document(OUTPUT)
    after_text = normalized_text(check)
    after_pages = page_structure(OUTPUT)

    if before_text != after_text:
        raise RuntimeError("Visible wording changed during formatting")
    if before_pages != after_pages:
        raise RuntimeError(f"Page or section structure changed: {before_pages} != {after_pages}")

    print(OUTPUT)
    print("normalized_text_match=true")
    print("page_structure_match=true")
    print(before_pages)


if __name__ == "__main__":
    main()
