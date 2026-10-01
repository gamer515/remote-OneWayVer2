from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


SOURCE = Path(r"C:\Users\유영선\Desktop\C211201_유영선.docx")
OUTPUT = Path(r"C:\remote-OneWayVer2\outputs\C211201_유영선_서식정리본.docx")
FONT = "맑은 고딕"


MAJOR_HEADINGS = {
    "과제 개요",
    "필요성 및 동기",
    "요소 기술",
    "상세 내용",
    "주간별/팀원별 과제 수행 계획",
}

SECTION_HEADINGS = {
    "AI 협업 개발 역량과 반복 가능한 제작 과정",
    "기존 게임과 서비스 분석",
    "반복 개발과 통합 검증",
    "공통",
    "기존 기술과 구현 자산의 재사용",
    "차별화기술",
    "중점기술",
    "요구 사항",
    "시스템 설계",
    "구현",
    "시험평가",
    "팀원별 역할",
    "전체 시스템 계층",
    "주요 자료구조",
    "주요 알고리즘",
    "전체 시스템 계층:",
    "씬 컨트롤:",
    "주요 자료구조(데이터:",
    "주요 알고리즘:",
}

SUBHEADINGS = {
    "기능적 요구사항",
    "기능적 요구사항:",
    "비기능적 요구사항",
    "비기능적 요구사항:",
}

TECH_TITLES = {
    "Unity 6.3 LTS와 C#",
    "Git과 GitHub Desktop",
    "Codex CLI MCP 플러그인과 skill",
    "Blender Pixilart와 Unity 에셋",
    "Api 기반 LLM Model",
}


def set_run_font(run, size: float, bold: bool | None = None) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), FONT)


def normalize_layout_whitespace(paragraph) -> None:
    for run in paragraph.runs:
        if "\t" in run.text or "\n" in run.text:
            run.text = run.text.replace("\t", " ").replace("\n", " ")


def clear_indents(paragraph) -> None:
    fmt = paragraph.paragraph_format
    fmt.left_indent = Cm(0)
    fmt.right_indent = Cm(0)
    fmt.first_line_indent = Cm(0)


def format_document_paragraph(paragraph) -> None:
    normalize_layout_whitespace(paragraph)
    text = paragraph.text.strip()
    fmt = paragraph.paragraph_format

    clear_indents(paragraph)
    fmt.keep_together = False
    fmt.keep_with_next = False
    fmt.widow_control = True

    if not text:
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.line_spacing = 1.0
        for run in paragraph.runs:
            set_run_font(run, 1)
        return

    if text in MAJOR_HEADINGS:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.space_before = Pt(11)
        fmt.space_after = Pt(5)
        fmt.line_spacing = 1.0
        fmt.keep_with_next = True
        for run in paragraph.runs:
            set_run_font(run, 14, True)
        return

    if text in SECTION_HEADINGS:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.space_before = Pt(8)
        fmt.space_after = Pt(3)
        fmt.line_spacing = 1.0
        fmt.keep_with_next = True
        for run in paragraph.runs:
            set_run_font(run, 11.5, True)
        return

    if text in SUBHEADINGS:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.left_indent = Cm(0.4)
        fmt.space_before = Pt(5)
        fmt.space_after = Pt(2)
        fmt.line_spacing = 1.0
        fmt.keep_with_next = True
        for run in paragraph.runs:
            set_run_font(run, 10.5, True)
        return

    if text in {"유영선:", "황준호:"}:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.left_indent = Cm(0.3)
        fmt.space_before = Pt(5)
        fmt.space_after = Pt(2)
        fmt.line_spacing = 1.0
        fmt.keep_with_next = True
        for run in paragraph.runs:
            set_run_font(run, 10.5, True)
        return

    if text in TECH_TITLES:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.left_indent = Cm(0.4)
        fmt.space_before = Pt(4)
        fmt.space_after = Pt(1)
        fmt.line_spacing = 1.0
        fmt.keep_with_next = True
        for run in paragraph.runs:
            set_run_font(run, 10.5, True)
        return

    if text.startswith("•"):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.left_indent = Cm(0.9)
        fmt.first_line_indent = Cm(-0.45)
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(2)
        fmt.line_spacing = 1.15
        for run in paragraph.runs:
            set_run_font(run, 10.5)
        return

    if re.match(r"^\d+\.\s*", text):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.left_indent = Cm(0.85)
        fmt.first_line_indent = Cm(-0.5)
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(2)
        fmt.line_spacing = 1.15
        for run in paragraph.runs:
            set_run_font(run, 10.5)
        return

    label_prefixes = (
        "역할:", "적용:", "구조:", "운영체제:", "개발 언어:",
        "게임 엔진과 IDE:", "데이터 도구:", "AI와 자동화 도구:",
        "그래픽 도구:", "협업 도구:",
    )
    if text.startswith(label_prefixes):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.left_indent = Cm(0.4)
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(2)
        fmt.line_spacing = 1.15
        for run in paragraph.runs:
            set_run_font(run, 10.5)
        return

    if text.startswith(("유영선:", "황준호:")):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.left_indent = Cm(0.4)
        fmt.space_before = Pt(2)
        fmt.space_after = Pt(3)
        fmt.line_spacing = 1.2
        for run in paragraph.runs:
            set_run_font(run, 10.5)
        return

    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if len(text) >= 45 else WD_ALIGN_PARAGRAPH.LEFT
    fmt.first_line_indent = Cm(0.6) if len(text) >= 45 else Cm(0)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(3)
    fmt.line_spacing = 1.2
    for run in paragraph.runs:
        set_run_font(run, 10.5)


def format_table(table, index: int) -> None:
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                fmt = paragraph.paragraph_format
                fmt.space_before = Pt(0)
                fmt.space_after = Pt(0)
                fmt.line_spacing = 1.0
                fmt.left_indent = Cm(0)
                fmt.right_indent = Cm(0)
                fmt.first_line_indent = Cm(0)

                if index == 0:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    size = 12 if r_idx == 0 else 10
                    bold = r_idx == 0 or c_idx == 0
                elif index == 1:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 and r_idx > 0 else WD_ALIGN_PARAGRAPH.CENTER
                    size = 8.5 if c_idx == 0 else 8
                    bold = r_idx == 0 or c_idx == 0
                else:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
                    size = 8
                    bold = c_idx == 0

                for run in paragraph.runs:
                    set_run_font(run, size, bold)

    if index == 1 and table.rows:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        repeat = tr_pr.find(qn("w:tblHeader"))
        if repeat is None:
            from docx.oxml import OxmlElement

            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            tr_pr.append(repeat)


def normalized_text(doc) -> str:
    items: list[str] = []
    items.extend(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                items.extend(p.text for p in cell.paragraphs)
    for section in doc.sections:
        items.extend(p.text for p in section.header.paragraphs)
        items.extend(p.text for p in section.footer.paragraphs)
    return re.sub(r"\s+", " ", "\n".join(items)).strip()


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(SOURCE)
    before = normalized_text(doc)

    for style_name in ("Normal", "Body Text", "List Paragraph"):
        style = doc.styles[style_name]
        style.font.name = FONT
        style.font.size = Pt(10.5)
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style._element.rPr.rFonts.set(qn("w:cs"), FONT)

    heading = doc.styles["Heading 1"]
    heading.font.name = FONT
    heading.font.size = Pt(14)
    heading.font.bold = True
    heading._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    heading._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    heading._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    heading._element.rPr.rFonts.set(qn("w:cs"), FONT)

    for paragraph in doc.paragraphs:
        format_document_paragraph(paragraph)

    for index, table in enumerate(doc.tables):
        format_table(table, index)

    if len(doc.sections) > 1:
        doc.sections[1].start_type = WD_SECTION.CONTINUOUS

    doc.save(OUTPUT)
    check = Document(OUTPUT)
    after = normalized_text(check)
    if before != after:
        raise RuntimeError("Visible wording changed during formatting")

    print(OUTPUT)
    print(f"paragraphs={len(check.paragraphs)} tables={len(check.tables)} sections={len(check.sections)}")
    print("normalized_text_match=true")


if __name__ == "__main__":
    main()
