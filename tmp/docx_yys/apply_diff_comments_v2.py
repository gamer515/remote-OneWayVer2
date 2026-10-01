from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


SOURCE = Path(r"C:\remote-OneWayVer2\outputs\C211201_유영선_보고서양식준수본.docx")
OUTPUT = Path(r"C:\remote-OneWayVer2\outputs\C211201_유영선_보고서양식준수본_v2.docx")

FONT_NAME = "바탕체"
BODY_SIZE = Pt(10)


def set_run_font(run, size=BODY_SIZE):
    run.font.name = FONT_NAME
    run.font.size = size
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)


def style_hanging_item(paragraph, left_cm=1.25, hanging_cm=0.45):
    fmt = paragraph.paragraph_format
    fmt.left_indent = Cm(left_cm)
    fmt.first_line_indent = Cm(-hanging_cm)
    fmt.space_after = Pt(1)
    fmt.line_spacing = 1.25


def style_indented_paragraph(paragraph, left_cm):
    fmt = paragraph.paragraph_format
    fmt.left_indent = Cm(left_cm)
    fmt.first_line_indent = Cm(0)
    fmt.space_after = Pt(1)
    fmt.line_spacing = 1.25


def replace_text(paragraph, text):
    paragraph.clear()
    run = paragraph.add_run(text)
    set_run_font(run)


def replace_with_lines(paragraph, lines):
    paragraph.clear()
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        set_run_font(run)
        if index < len(lines) - 1:
            run.add_break()
    fmt = paragraph.paragraph_format
    fmt.left_indent = Cm(0.8)
    fmt.first_line_indent = Cm(0)
    fmt.space_after = Pt(1)
    fmt.line_spacing = 1.25


def structural_signature(path: Path):
    doc = Document(path)
    sections = []
    for s in doc.sections:
        sections.append(
            (
                s.start_type,
                s.page_width,
                s.page_height,
                s.top_margin,
                s.bottom_margin,
                s.left_margin,
                s.right_margin,
            )
        )
    with zipfile.ZipFile(path) as zf:
        xml = zf.read("word/document.xml")
    return {
        "sections": sections,
        "page_breaks": xml.count(b'w:type="page"'),
        "last_rendered_page_breaks": xml.count(b"w:lastRenderedPageBreak"),
        "sect_prs": xml.count(b"<w:sectPr"),
    }


before = structural_signature(SOURCE)
doc = Document(SOURCE)

# 모든 글머리표 문단에 양식에 맞는 내어쓰기 적용.
for paragraph in doc.paragraphs:
    if paragraph.text.strip().startswith("•"):
        style_hanging_item(paragraph)

# API 기반 LLM 항목은 기존 문구를 유지하면서 계층형 들여쓰기만 적용.
for index, paragraph in enumerate(doc.paragraphs):
    if paragraph.text.strip() == "Api 기반 LLM Model":
        style_indented_paragraph(paragraph, 0.8)
        if index + 1 < len(doc.paragraphs):
            style_indented_paragraph(doc.paragraphs[index + 1], 1.2)
        break
else:
    raise RuntimeError("'Api 기반 LLM Model' paragraph not found")

# 탄막 패턴 선택 하위 항목은 상위 숫자 번호를 유지하고 A-D 알파벳 번호를 부여.
alphabet_items = {
    "순차 실행:": "A.",
    "무작위 실행:": "B.",
    "가중치 실행:": "C.",
    "조건부 실행:": "D.",
}
found_alpha = set()
for paragraph in doc.paragraphs:
    stripped = paragraph.text.strip()
    for prefix, marker in alphabet_items.items():
        if stripped.startswith(prefix):
            replace_text(paragraph, f"{marker} {stripped}")
            style_hanging_item(paragraph, left_cm=1.35, hanging_cm=0.55)
            found_alpha.add(prefix)
            break
if found_alpha != set(alphabet_items):
    raise RuntimeError(f"Missing algorithm items: {set(alphabet_items) - found_alpha}")

# AI/자동화 도구 담당을 이름별로 줄바꿈하여 구분.
for paragraph in doc.paragraphs:
    if paragraph.text.strip().startswith("AI와 자동화 도구:"):
        replace_with_lines(
            paragraph,
            [
                "AI와 자동화 도구:",
                "유영선: 한국어 중심 LLM의 파인튜닝, 경량화 및 게임 내장형 추론 환경을 구성합니다.",
                "황준호: 외부 LLM API 통신을 이용합니다.",
                "개발 과정에서는 Codex, CLI, MCP, 플러그인과 skill을 활용합니다.",
            ],
        )
        break
else:
    raise RuntimeError("AI와 자동화 도구 paragraph not found")

# 유영선 일정의 마지막 항목을 15주차까지 계획(P)으로 확장.
schedule_updated = False
for table in doc.tables:
    for row in table.rows:
        row_text = " ".join(cell.text.strip() for cell in row.cells)
        if "통합 성능 사용성 테스트" in row_text:
            target = row.cells[-1]
            target.text = "P"
            for paragraph in target.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    set_run_font(run)
            schedule_updated = True
            break
    if schedule_updated:
        break
if not schedule_updated:
    raise RuntimeError("Schedule row not found")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)

after = structural_signature(OUTPUT)
if before != after:
    raise RuntimeError(f"Page/section structure changed: before={before}, after={after}")

print(OUTPUT)
print(before)
